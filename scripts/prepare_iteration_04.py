"""Prepare versioned simulation inputs from reviewed sources; never edit raw/OD.

Optional preparation dependencies: openpyxl and pypdf. The runtime engine is stdlib.
Public report observations were visually/text reviewed; assert the exact series.
"""
import argparse
import csv
import hashlib
import json
import re
import sys
from copy import deepcopy
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from pilog_cfr_dss.models.iteration_04 import UNITS
from pilog_cfr_dss.calculations.fuel import port_diesel_mt
from pilog_cfr_dss.orchestration.iteration_04 import evaluate, fixed_freight_record


def prepare_case(source: Path, output: Path) -> None:
    """Prepare an existing structured case without workbook/PDF dependencies.

    A new destination is required. Existing fixed prices are preserved; to price
    a deliberately NEW offer, remove freight_tender in a copy before preparation.
    """
    report = evaluate(source)
    record = fixed_freight_record(report)
    context = json.loads(report["_captured_inputs"]["case_fixture.json"])
    name = context["voyage_dataset"]
    targets = [output / source.name, output / name]
    if any(p.exists() for p in targets):
        raise ValueError("Destination input exists; choose a new --output-dir.")
    context["freight_tender"] = record
    output.mkdir(parents=True, exist_ok=True)
    targets[1].write_bytes(report["_captured_inputs"][name])
    targets[0].write_text(json.dumps(context, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"prepared": str(targets[0]), "freight_tender": record, "commercial_approval": False}, indent=2))


def build(output: Path) -> None:
    from openpyxl import load_workbook
    from pypdf import PdfReader
    files = [output / "i4_voyage.csv", output / "i4_context.json", output / "i4_no_go.json"]
    if any(p.exists() for p in files):
        raise ValueError("I4 inputs already exist; choose a new --output-dir instead of overwriting experiments.")
    raw = ROOT / "data/raw/Voyage Estimate.xlsx"
    od_path = ROOT / "docs/it_4/Simulasi Tarif Internasional_Mei 2025.xlsx"
    register = ROOT / "data/raw/Voyage_Estimate_all_sheets_variable_value_units.xlsx"
    sources = {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in [raw, od_path, register]}
    w = load_workbook(raw, read_only=True, data_only=True)
    formulas = load_workbook(raw, read_only=True, data_only=False)
    od = load_workbook(od_path, read_only=True, data_only=True)
    aq = w["AQABA"]
    od_input = od["Input - OD Intl"]
    ref = "data/raw/Voyage Estimate.xlsx :: "
    observations = {}
    report_paths = {
        "argus": ROOT.parent / "references/Market Reports/Argus Nitrogen (2026-09-10).pdf",
        "fertecon": ROOT.parent / "references/Market Reports/Fertecon Nitrogen Market Report 2026-09-10.pdf",
    }
    patterns = {
        "argus": (1, r"Southeast Asia\s+(\d+)-(\d+)\s+405-420"),
        "fertecon": (0, r"FOB SE Asia urea \(granular\)\s+FGUAA04\s+[–-]\s+(\d+\.\d+)\s*[–-]\s*(\d+\.\d+)"),
    }
    for publisher, path in report_paths.items():
        page, pattern = patterns[publisher]
        content = PdfReader(path).pages[page].extract_text()
        # Argus row is specifically under granular FOB bulk, not prilled/CFR.
        if publisher == "argus":
            content = content.split("Granular urea - fob bulk", 1)[1]
        match = re.search(pattern, content)
        if match is None:
            raise ValueError("REPORT_SERIES_NOT_FOUND:" + publisher)
        observations[publisher] = {
            "low": match[1], "high": match[2], "unit": "USD/MT", "basis": "FOB SEA granular",
            "series": "Granular urea FOB bulk Southeast Asia" if publisher == "argus" else "FOB SE Asia urea granular FGUAA04",
            "observation_date": "2026-09-10", "source_status": "observed_report",
            "source_reference": str(path.relative_to(ROOT.parent)) + f" :: PDF page {page + 1}",
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        }
    baltic = "https://www.balticexchange.com/en/data-services/market-information0/dry-services.html"
    rows = []

    def add(key, value, status, source, note=""):
        rows.append({"key": key, "value": str(value), "unit": UNITS[key], "source_status": status,
                     "source_reference": source, "note": note,
                     "role": "reference_unused_gross_dwt_only" if key in {"bunkers_allowance_mt", "constants_allowance_mt"} else "input"})

    for key, value in {"cargo_description": "UREA GRANULAR BULK", "cargo_category": "dry_bulk",
                       "loading_port": "Bontang", "discharging_port": "Mundra", "vessel_class": "Handysize",
                       "vessel_mode": "owned", "voyage_boundary": "laden_only", "cargo_mt": "30000"}.items():
        add(key, value, "confirmed_simulation_case", "Mentor/user I4 plan and D-040/D-042")
    add("distance_nm", od_input["H11"].value, "workbook_reference", "docs/it_4/Simulasi Tarif Internasional_Mei 2025.xlsx :: Input - OD Intl!H11", "One-way reference, not independently verified navigation distance")
    for key, value in {"dwt_mt": "38200", "speed_knots": "12", "sea_fo_mt_day": "18", "sea_do_mt_day": "0.1"}.items():
        add(key, value, "benchmark_assumption", baltic + " :: Handysize 38 at 12 knots", "Sea consumption includes ME+AE; not actual PILOG vessel. IFO/MDO prices below are proxies.")
    for key, cell in {"bunkers_allowance_mt": "D18", "constants_allowance_mt": "D19",
                      "fo_price_usd_mt": "G24", "do_price_usd_mt": "G28", "running_usd_day": "F44"}.items():
        add(key, aq[cell].value, "inherited_simulation_proxy", ref + "AQABA!" + cell,
            "Historical template input carried to new case; not current Handysize/route observation")
    port_consumption = port_diesel_mt(Decimal(1), Decimal(2000), Decimal("0.7"), Decimal("0.1"))
    add("port_do_mt_day", port_consumption, "derived_simulation_proxy", ref + "AQABA!J6 = 0.7*0.1*2000*24/1000", "Port engine proxy separate from Baltic sea-only benchmark; not a full Baltic vessel profile")
    for key, cell in [("loading_rate_mt_day", "D12"), ("discharge_rate_mt_day", "F12")]:
        formula = formulas["Sby"][cell].value
        match = re.fullmatch(r"=D20/(\d+(?:\.\d+)?)", str(formula))
        if match is None:
            raise ValueError("UNSUPPORTED_HANDLING_FORMULA:" + cell)
        add(key, match[1], "inherited_simulation_proxy", ref + "Sby!" + cell + " embedded divisor: " + formula,
            "Reviewed Sby structure carried to synthetic route; not actual terminal rate")
    add("waiting_days", "0", "simulation_only_assumption", "I4 fixture", "Explicit baseline no extra waiting; stress adds port delay")
    add("port_other_total_usd", aq["J33"].value, "simulation_only_assumption", ref + "AQABA!J33; mentor aggregate interpretation from Sby audit", "Use once as TOTAL port/other for this simulation, excludes bunker/running. Do not copy Sby 250000 additionally or claim actual Mundra tariff.")
    for key, cell in [("freight_commission_fraction", "H41"), ("broker_fraction", "L46"), ("additional_commission_fraction", "L48")]:
        formula = formulas["AQABA"][cell].value
        percentages = re.findall(r"(\d+(?:\.\d+)?)%", str(formula))
        if len(percentages) != 1:
            raise ValueError("UNSUPPORTED_COMMISSION_FORMULA:" + cell)
        fraction = Decimal(percentages[0]) / Decimal(100)
        add(key, fraction, "simulation_only_assumption", ref + "AQABA!" + cell + " embedded percentage: " + formula,
            "Percentage extracted from formula, NOT cached USD amount. Cumulative gross-revenue basis is a simulation assumption, not verified contract.")
    if set(r["key"] for r in rows) != set(UNITS):
        raise ValueError("PREPARATION_SCHEMA_MISMATCH")
    bids = []
    alphas = {1: [-10, -8, -5], 2: [2, 3, 1], 3: [12, 8, 6]}
    for rd, values in alphas.items():
        for seq, alpha in enumerate(values, 1):
            bids.append({"bid_id": f"R{rd}-B{seq}", "buyer_id": f"BUYER-{seq}", "round": rd, "sequence": seq,
                         "alpha_usd_mt": str(alpha), "status": "active", "source_status": "simulation_only_synthetic"})
    captured = datetime.now(timezone.utc).isoformat()
    context = {
        "schema_version": "I4.1", "mode": "simulation", "case_id": "I4_GO_BONTANG_MUNDRA",
        "voyage_dataset": "i4_voyage.csv", "evaluation_cutoff": "2026-09-10T12:00:00+00:00",
        "market_index": observations, "oe_hps": {"value": "420", "unit": "USD/MT", "source_status": "simulation_only_synthetic", "source_reference": "I4 PI-floor fixture, not PI actual OE/HPS"},
        "pricing": {"basis": "tce_net_voyage_before_running", "quote_commission_basis": "simulation_assumed_net",
                    "target_tce": {"value": str(aq["L38"].value), "unit": "USD/day", "source_status": "simulation_only_proxy",
                                   "source_reference": ref + "AQABA!L38", "note": "Historical hire input reinterpreted ONLY as simulation target TCE; not actual Fearnpulse quote or owned-vessel hire expense"}},
        "market_sources": {"fearnpulse": {"url": "https://fearnpulse.com/#2", "capture_date": captured,
                                            "status": "dynamic_page_no_verified_numeric_quote", "fallback": "labelled AQABA L38 target-TCE proxy, not a Fearnpulse observation"},
                           "baltic": {"url": baltic, "capture_date": captured, "profile": "Handysize 38 at 12kn; benchmark sea profile only"}},
        "od_reference": {"value": str(od["Output - Freight"]["F12"].value), "unit": "USD/MT", "source_status": "observed_workbook_reference",
                         "source_reference": "docs/it_4/Simulasi Tarif Internasional_Mei 2025.xlsx :: Output - Freight!F12 / Process - Intl!AG65",
                         "sha256": sources[str(od_path.relative_to(ROOT))], "quote_basis": "one_way_shipment", "cost_coverage": "includes_return",
                         "loading_port": "Bontang", "discharging_port": "Mundra", "cargo_mt": "30000", "vessel_class": "Handysize",
                         "source_period": "filename May2025; output also labels November2024; retained as supplied, not current-market quote"},
        "vessel_policy": {"Handysize": ["20000", "40000"], "Supramax": ["50000", "60000"]},
        "bids": bids,
        "scenarios": [{"scenario_id": "fuel_up", "label": "BBM naik 25%", "fuel_price_increase_fraction": "0.25", "extra_port_days": "0", "source_status": "simulation_only_assumption"},
                      {"scenario_id": "port_delay", "label": "Delay pelabuhan 30 hari", "fuel_price_increase_fraction": "0", "extra_port_days": "30", "source_status": "simulation_only_assumption"}],
        "source_manifest": sources,
        "limitations": ["Urea/Handysize/rute baru adalah kasus simulasi, bukan voyage aktual template.",
                        "Sistem menghitung biaya satu arah; biaya balik berada di luar model. OD tarif satu arah memakai biaya internal termasuk balik.",
                        "3.906 NM acuan workbook; belum diverifikasi independen.",
                        "Port/other dan profil pelabuhan adalah proxy berlabel; bukan tarif/kapal aktual.",
                        "Fearnpulse belum menghasilkan quote numerik terverifikasi; target TCE memakai proxy historis AQABA, bukan angka Fearnpulse.",
                        "Laycan/terminal/DWCC/availability belum dibuktikan oleh screen DWT."]}
    w.close(); formulas.close(); od.close()
    # Freeze source identity across preparation, not just cached values.
    for name, digest in sources.items():
        if hashlib.sha256((ROOT / name).read_bytes()).hexdigest() != digest:
            raise ValueError("SOURCE_CHANGED_DURING_PREPARATION:" + name)
    output.mkdir(parents=True, exist_ok=True)
    with files[0].open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)
    # Intentional pricing preparation, before any tender evaluation. No hardcoded rate.
    files[1].write_text(json.dumps(context, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    context["freight_tender"] = fixed_freight_record(evaluate(files[1]))
    no_go = deepcopy(context)
    no_go["case_id"] = "I4_NO_GO_TENDER_FLOOR"
    for bid in no_go["bids"]:
        bid["alpha_usd_mt"] = str(-20 + bid["round"] + bid["sequence"])
    for path, payload in [(files[1], context), (files[2], no_go)]:
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"created": [str(p) for p in files], "source_hashes_unchanged": True}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Siapkan input I4 dari sumber reviewed tanpa overwrite")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data/processed")
    parser.add_argument("--from-case", type=Path, help="Prepare a copy of structured inputs, no workbook/report reads")
    args = parser.parse_args()
    if args.from_case:
        prepare_case(args.from_case, args.output_dir)
    else:
        build(args.output_dir)
