"""Build the educational Iteration 2 variable map from two checked CSV inputs.

The workbook explains source fields. It never performs business calculations.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATASET = ROOT / "data" / "processed" / "dataset_ready.csv"
CATALOG = ROOT / "data" / "processed" / "variable_catalog.csv"
OUTPUT = ROOT / "outputs" / "learning" / "iteration_02_variable_map.xlsx"

GROUPS = {
    "case_and_commercial": "Kasus Komersial",
    "route_and_timing": "Rute Waktu",
    "vessel_and_capacity": "Kapal Kapasitas",
    "fuel_and_consumption": "BBM Konsumsi",
    "port_operations": "Operasi Pelabuhan",
    "cost_commission_revenue": "Biaya Pendapatan",
    "benchmark_and_market": "Benchmark Pasar",
}

HEADERS = [
    "No", "Variabel sumber", "Contoh nilai", "Kunci kanonis", "Definisi",
    "Tipe", "Satuan", "Mata uang", "Per rute?", "Per skenario?",
    "Peran", "Modul", "Pemakaian I1", "Pemakaian I2",
    "Status sumber", "Ketergantungan / pertanyaan",
]

NAVY = "17324D"
BLUE = "EAF2F8"
AMBER = "FFF1CC"
INK = "203247"
GRAY = "66788A"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def slug(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", label.strip().lower()).strip("_")


def example_value(raw: str):
    try:
        return float(Decimal(raw))
    except (InvalidOperation, ValueError):
        return raw


def joined_rows() -> list[dict]:
    source = read_csv(DATASET)
    catalog = read_csv(CATALOG)
    if not source or len(source) != len(catalog):
        raise ValueError("Dataset and catalog row counts must match")
    if any(slug(row["Variable"]) == "voyage_expense_allocation_quantity_mt" for row in source):
        raise ValueError("Obsolete 46,000 MT allocation field is not part of the active dataset")
    joined = []
    for source_row, spec in zip(source, catalog, strict=True):
        if source_row["No"] != spec["no"]:
            raise ValueError(f"Row number mismatch: {source_row['No']}")
        if slug(source_row["Variable"]) != spec["canonical_key"]:
            raise ValueError(f"Variable mismatch at row {source_row['No']}")
        if spec["group"] not in GROUPS:
            raise ValueError(f"Unknown group at row {source_row['No']}")
        joined.append({**spec, "raw_label": source_row["Variable"], "raw_value": source_row["Value"]})
    if len({row["canonical_key"] for row in joined}) != len(joined):
        raise ValueError("Duplicate canonical key")
    return joined


def prepare_sheet(sheet, title: str, subtitle: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Aptos", size=15, bold=True, color=NAVY)
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name="Aptos", size=10, italic=True, color=GRAY)
    sheet.row_dimensions[1].height = 27
    sheet.row_dimensions[2].height = 23


def style_table(sheet, last_row: int, last_col: int) -> None:
    header_fill = PatternFill("solid", fgColor=NAVY)
    thin = Side(style="thin", color="DCE5EC")
    for cell in sheet[4]:
        if cell.column > last_col:
            break
        cell.fill = header_fill
        cell.font = Font(name="Aptos", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    sheet.row_dimensions[4].height = 33
    for row_idx in range(5, last_row + 1):
        sheet.row_dimensions[row_idx].height = 45
        for col_idx in range(1, last_col + 1):
            cell = sheet.cell(row_idx, col_idx)
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F6F9FB")
            if col_idx == last_col and cell.value not in (None, "none"):
                cell.fill = PatternFill("solid", fgColor=AMBER)
        sheet.cell(row_idx, 1).alignment = Alignment(horizontal="center", vertical="top")
        sheet.cell(row_idx, 3).number_format = "#,##0.####"
    widths = [6, 35, 18, 39, 62, 12, 15, 13, 13, 15, 15, 17, 23, 29, 22, 34]
    for idx, width in enumerate(widths[:last_col], start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "C5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(last_col)}{last_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1


def map_sheet(workbook: Workbook, name: str, rows: list[dict], subtitle: str) -> None:
    sheet = workbook.create_sheet(name)
    prepare_sheet(sheet, name, subtitle)
    sheet.append([])
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([
            int(row["no"]), row["raw_label"], example_value(row["raw_value"]),
            row["canonical_key"], row["definition"], row["data_type"],
            row["unit"], row["currency"], row["route_dependent"],
            row["scenario_dependent"], row["role"], row["module"],
            row["current_use"], row["intended_use"], row["source_status"],
            row["missing_dependency"],
        ])
    style_table(sheet, 4 + len(rows), len(HEADERS))


def summary_sheet(workbook: Workbook, rows: list[dict]) -> None:
    sheet = workbook.active
    sheet.title = "Ringkasan"
    prepare_sheet(sheet, "Peta variabel Iteration 2", "Satu kasus simulasi; contoh nilai bukan tarif atau keputusan produksi")
    sheet["A4"] = "Kelompok"
    sheet["B4"] = "Jumlah variabel"
    counts = Counter(row["group"] for row in rows)
    for offset, (key, label) in enumerate(GROUPS.items(), start=5):
        sheet.cell(offset, 1, label)
        sheet.cell(offset, 2, counts[key])
    sheet.cell(12, 1, "Total")
    sheet.cell(12, 2, len(rows))
    sheet["A15"] = "Cara membaca"
    sheet["A16"] = "Contoh nilai berasal dari register eksperimen. Lihat kolom status sumber dan ketergantungan sebelum memakai angka."
    sheet["A17"] = "Per rute berarti nilainya harus mengikuti kasus yang dipilih; per skenario berarti dapat berubah menurut kapal atau opsi."
    sheet["A18"] = "Workbook ini menjelaskan data. Perhitungan resmi dilakukan oleh fungsi Python dan diuji terpisah."
    sheet["A20"] = "Sumber"
    sheet["A21"] = "Dataset eksperimen 61 variabel dan katalog variabel Iteration 2 revisi."
    sheet.column_dimensions["A"].width = 95
    sheet.column_dimensions["B"].width = 24
    for row_idx in (4, 15, 20):
        for cell in sheet[row_idx][:2]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Aptos", color="FFFFFF", bold=True)
    for row_idx in range(5, 13):
        sheet.cell(row_idx, 1).font = Font(name="Aptos", color=INK)
        sheet.cell(row_idx, 2).font = Font(name="Aptos", color=INK, bold=row_idx == 12)
    for row_idx in (16, 17, 18, 21):
        sheet.cell(row_idx, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row_idx].height = 33


def questions_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Pertanyaan Mentor")
    prepare_sheet(sheet, "Pertanyaan definisi", "Jawaban mentor diperlukan sebelum angka ekonomi menjadi dasar rekomendasi")
    questions = [
        ("Q01", "Freight rate", "Tarif jual PILOG per MT; sumber dan tanggal quote per kasus masih perlu dicatat.", "Definisi dikonfirmasi mentor"),
        ("Q02", "Kuantitas", "45.000 MT untuk semua hitungan; field alokasi 46.000 MT pada data lama salah dan dikeluarkan.", "Dikoreksi mentor"),
        ("Q03", "Other voyage costs", "USD 200.000 terpisah dari bunker dan running cost; rincian port/agency masih perlu diperiksa.", "Sebagian dikonfirmasi"),
        ("Q04", "Biaya kapal", "Simulasi ini memakai kapal PILOG dan running cost; hubungan dengan daily cost serta charter ditunda.", "Asumsi simulasi"),
        ("Q05", "Komisi", "Ketiga rate dihitung; apakah semuanya berbasis gross freight revenue?", "Basis masih perlu konfirmasi"),
        ("Q06", "Kapal", "Screen revisi ini hanya DWT; bunker/constants, draft, terminal, dan availability nyata belum diuji.", "Batas simulasi"),
        ("Q07", "Objective", "Kontribusi freight positif adalah sinyal review; minimum margin komersial ditunda.", "Perlu kebijakan nanti"),
        ("Q08", "Rute", "Apa sumber dan tanggal jarak laden/ballast serta biaya pelabuhan untuk kasus terpilih?", "Menunggu konfirmasi"),
    ]
    sheet.append([])
    sheet.append(["ID", "Topik", "Pertanyaan", "Status"])
    for item in questions:
        sheet.append(list(item))
    style_table(sheet, 4 + len(questions), 4)
    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 100
    sheet.column_dimensions["D"].width = 26


def main() -> None:
    rows = joined_rows()
    workbook = Workbook()
    summary_sheet(workbook, rows)
    map_sheet(workbook, "Semua Variabel", rows, "Filter kolom untuk melihat peran, satuan, dan ketergantungan setiap baris")
    for key, label in GROUPS.items():
        group_rows = [row for row in rows if row["group"] == key]
        map_sheet(workbook, label, group_rows, "Contoh nilai berasal dari satu dataset eksperimen")
    questions_sheet(workbook)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(f"Created {OUTPUT} with {len(rows)} mapped variables and {len(workbook.sheetnames)} sheets")


if __name__ == "__main__":
    main()
