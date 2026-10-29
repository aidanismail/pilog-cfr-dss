from copy import deepcopy
from decimal import Decimal as D
import hashlib
from pathlib import Path
import shutil
import tempfile
import unittest
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from pilog_cfr_dss.application.market_case import (
    catalog, voyage_catalog, new_market_case, apply_voyage, apply_profile,
    seed_demo_bids, demo_case, compare_profiles)
from pilog_cfr_dss.sources.fearnpulse import latest_saved
from pilog_cfr_dss.sources.voyage_estimates import read_voyages
from pilog_cfr_dss.calculations.voyage_04 import port_expenses_total
from pilog_cfr_dss.orchestration.iteration_04 import evaluate

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data/raw/Voyage Estimate.xlsx"
OD = ROOT / "docs/it_4/Simulasi Tarif Internasional_Mei 2025.xlsx"


class VoyageSelectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.workspace = Path(self.temp.name)
        self.drafts, self.runs = self.workspace / "drafts", self.workspace / "runs"
        self.snapshot = latest_saved(ROOT)

    def test_four_routes_with_distinct_cargo_and_od_only_distances(self):
        voyages = voyage_catalog(ROOT)["voyages"]
        self.assertEqual({v["fields"]["discharging_port"]: (v["fields"]["cargo_mt"], v["distance_reference"]["value"]) for v in voyages},
                         {"Mundra": ("30000", "3906"), "Davao": ("28000", "697"),
                          "Bangkok": ("32000", "1751"), "Hai Phong": ("29000", "1577")})
        self.assertTrue(all(v["source_status"] == "simulation_only_workbook_case" for v in voyages))
        self.assertTrue(all("asumsi demo" in v["note"] for v in voyages))

    def test_all_sixteen_pairs_reach_correct_economics_and_offline_replay(self):
        for voyage in voyage_catalog(ROOT)["voyages"]:
            for profile in catalog(ROOT)["profiles"]:
                with self.subTest(voyage=voyage["voyage_id"], profile=profile["profile_id"]):
                    case = new_market_case(ROOT, self.snapshot, profile["profile_id"], voyage["voyage_id"])
                    report = case.prepare(self.drafts)
                    self.assertEqual(report["case"]["cargo_mt"], voyage["fields"]["cargo_mt"])
                    self.assertEqual(report["case"]["discharging_port"], voyage["fields"]["discharging_port"])
                    self.assertEqual(report["case"]["speed_knots"], profile["fields"]["speed_knots"])
                    self.assertEqual(report["case"]["dwt_mt"], profile["fields"]["dwt_mt"])
                    self.assertEqual(report["route_distance"]["value"], voyage["distance_reference"]["value"])
                    self.assertEqual(report["baseline"]["port_other_usd"], voyage["fields"]["port_other_total_usd"])
                    case.fix_freight(self.drafts)
                    seed_demo_bids(case, ROOT)
                    result = case.evaluate(self.drafts, self.runs)
                    self.assertEqual(result["recommendation"]["action"], "GO")
                    self.assertEqual(result["input_lineage"]["context"]["voyage_estimate"]["voyage_id"], voyage["voyage_id"])
                    replay = evaluate(case.run_dir / "input_snapshot/case_fixture.json")
                    self.assertEqual(result["baseline"], replay["baseline"])

    def test_profile_switch_does_not_switch_voyage_and_reverse_is_independent(self):
        case = new_market_case(ROOT, self.snapshot, voyage_id="bontang_davao")
        voyage_before = deepcopy(case.context["voyage_estimate"])
        profiles = catalog(ROOT)
        apply_profile(case, profiles["profiles"][-1], profiles["policy"])
        self.assertEqual(case.context["voyage_estimate"], voyage_before)
        profile_before = deepcopy(case.context["vessel_profile"])
        apply_voyage(case, voyage_catalog(ROOT)["voyages"][2], ROOT)
        self.assertEqual(case.context["vessel_profile"], profile_before)
        self.assertEqual(case.values()["discharging_port"], "Bangkok")
        self.assertEqual(case.values()["waiting_days"], "1")

    def test_cached_excel_results_match_engine_for_the_profile_on_each_sheet(self):
        workbook = load_workbook(RAW, read_only=True, data_only=True)
        self.addCleanup(workbook.close)
        profiles = {p["sheet"]:p for p in catalog(ROOT)["profiles"]}
        for voyage in voyage_catalog(ROOT)["voyages"]:
            with self.subTest(sheet=voyage["sheet"]):
                profile = profiles[voyage["sheet"]]
                case = new_market_case(ROOT, self.snapshot, profile["profile_id"], voyage["voyage_id"])
                sheet = workbook[voyage["sheet"]]
                # Excel retains its saved market values; a later source refresh
                # must not invalidate this workbook-vs-engine reconciliation.
                case.apply_inputs({"fo_price_usd_mt":str(sheet["G24"].value),
                                   "do_price_usd_mt":str(sheet["G28"].value)},
                                  target_tce=str(sheet["K12"].value), observations={},
                                  oe_hps=case.context["oe_hps"]["value"])
                prepared = case.prepare(self.drafts)
                self.assertEqual(D(str(sheet["D41"].value)), D(prepared["pricing"]["computed_freight_usd_mt"]))
                self.assertLess(abs(D(str(sheet["H13"].value))-D(prepared["baseline"]["total_days"])),D("0.000000001"))
                self.assertLess(abs(D(str(sheet["J38"].value))-D(prepared["baseline"]["voyage_expenses_usd"])),D("0.000001"))

    def test_fixed_case_cannot_change_route_and_revised_case_discards_bids(self):
        case = new_market_case(ROOT, self.snapshot)
        case.prepare(self.drafts); case.fix_freight(self.drafts); seed_demo_bids(case, ROOT)
        old = case.fingerprint
        voyage = voyage_catalog(ROOT)["voyages"][1]
        with self.assertRaisesRegex(ValueError, "BASIS_FIXED"):
            apply_voyage(case, voyage, ROOT)
        self.assertEqual(case.fingerprint, old)
        revised = case.revise(); apply_voyage(revised, voyage, ROOT)
        self.assertFalse(revised.fixed)
        self.assertEqual(revised.context["bids"], [])
        self.assertEqual(revised.values()["discharging_port"], "Davao")

    def test_both_no_go_reasons_work_on_each_route(self):
        for voyage in voyage_catalog(ROOT)["voyages"]:
            for kind, gate_b in (("no-go-floor", "NOT_EVALUATED"), ("no-go-economics", "FAIL")):
                with self.subTest(voyage=voyage["voyage_id"], kind=kind):
                    case = demo_case(ROOT, self.snapshot, self.drafts, kind, voyage_id=voyage["voyage_id"])
                    result = case.evaluate(self.drafts, self.runs)
                    self.assertEqual(result["recommendation"]["action"], "NO_GO")
                    self.assertEqual(result["recommendation"]["gate_b"], gate_b)
                    self.assertEqual(result["case"]["discharging_port"], voyage["fields"]["discharging_port"])

    def test_comparison_keeps_route_and_quantity(self):
        case = new_market_case(ROOT, self.snapshot, voyage_id="bontang_bangkok")
        fingerprint = case.fingerprint
        rows = compare_profiles(case, ROOT, self.drafts)
        self.assertEqual(len(rows), 4)
        self.assertEqual(case.fingerprint, fingerprint)
        self.assertTrue(all(r["baseline"]["port_other_usd"] == "120000" for r in rows))

    def test_workbook_changes_invalidate_catalog_and_renaming_preserves_profile_ids(self):
        root = self.workspace / "repo"
        (root / "data/raw").mkdir(parents=True)
        (root / "docs/it_4").mkdir(parents=True)
        path = root / "data/raw/Voyage Estimate.xlsx"
        shutil.copy2(RAW, path); shutil.copy2(OD, root / "docs/it_4" / OD.name)
        before = voyage_catalog(root)
        workbook = load_workbook(path)
        workbook["Bontang-Davao"]["D20"] = 28500
        workbook["Bontang-Davao"].title = "Davao revised"
        workbook.save(path); workbook.close()
        after = voyage_catalog(root)
        self.assertNotEqual(before["source_sha256"], after["source_sha256"])
        self.assertEqual(after["voyages"][1]["fields"]["cargo_mt"], "28500")
        self.assertEqual(catalog(root)["profiles"][1]["profile_id"], "ve_aqaba")
        # Returned records cannot mutate the process-wide immutable-source cache.
        after["voyages"][1]["fields"]["cargo_mt"] = "999"
        self.assertEqual(voyage_catalog(root)["voyages"][1]["fields"]["cargo_mt"], "28500")

    def test_ambiguous_quantity_od_mismatch_and_missing_cost_are_not_silent_defaults(self):
        for cell, value, error in (("B41", 999, "CARGO_REVENUE"), ("F3", 999, "OD_DISTANCE"),
                                   ("I33", None, "I33"), ("D12", "=D20/SUM(1,2)", "HANDLING"),
                                   ("H10", 1, "RETURN_LEG")):
            with self.subTest(cell=cell):
                path = self.workspace / "bad.xlsx"
                workbook = load_workbook(RAW); workbook.worksheets[0][cell] = value
                workbook.save(path); workbook.close()
                with self.assertRaisesRegex(ValueError, error):
                    read_voyages(path, OD)

    def test_port_total_excludes_fuel_and_running_and_rejects_invalid_components(self):
        self.assertEqual(port_expenses_total([D(45000),D(40000),D(0),D(0),D(5000)]), D(90000))
        for values in ([],[D(-1)],[D("NaN")]):
            with self.assertRaises(ValueError): port_expenses_total(values)

    def test_original_formatting_parts_and_cell_styles_unchanged(self):
        original = ROOT / "docs/it_4/Voyage Estimate I4.xlsx"
        with ZipFile(original) as a, ZipFile(RAW) as b:
            self.assertEqual(a.namelist(), b.namelist())
            mutable = {"xl/workbook.xml"} | {f"xl/worksheets/sheet{i}.xml" for i in range(1,5)}
            for name in a.namelist():
                if name not in mutable: self.assertEqual(a.read(name),b.read(name),name)
            ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            for i in range(1,5):
                aroot=ET.fromstring(a.read(f"xl/worksheets/sheet{i}.xml"))
                broot=ET.fromstring(b.read(f"xl/worksheets/sheet{i}.xml"))
                for element in ["cols","mergeCells","sheetFormatPr","pageMargins","sheetViews"]:
                    self.assertEqual(ET.tostring(aroot.find("s:"+element,ns)),ET.tostring(broot.find("s:"+element,ns)))
                acells={c.attrib["r"]:c.attrib.get("s") for c in aroot.findall(".//s:c",ns)}
                bcells={c.attrib["r"]:c.attrib.get("s") for c in broot.findall(".//s:c",ns)}
                self.assertTrue(all(bcells.get(k)==v for k,v in acells.items()))
