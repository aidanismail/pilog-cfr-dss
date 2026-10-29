"""Read OD route distances only; never import the workbook's freight prices."""
import hashlib
from copy import deepcopy
from functools import lru_cache
from io import BytesIO
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

SHEET = "Input - OD Intl"
PRODUCT_ALIASES = {"urea granular bulk": "urea"}


def payload_bytes(payload: dict) -> bytes:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _label(value) -> str:
    return " ".join(str(value).strip().casefold().split())


def resolve_distance(payload: dict, origin: str, destination: str, product: str) -> dict:
    """Resolve one product/origin/destination intersection in captured raw cells."""
    if payload.get("sheet") != SHEET or payload.get("unit") != "NM":
        raise ValueError("INVALID_OD_DISTANCE_UNIT_OR_SHEET")
    cells = payload["cells"]
    columns = [col for col in "FGHIJ" if _label(cells.get(col + "5", "")) == _label(origin)]
    if len(columns) != 1:
        raise ValueError("OD_DISTANCE_ORIGIN_NOT_UNIQUE")
    product = PRODUCT_ALIASES.get(_label(product), _label(product))
    current_product, product_cell = "", None
    matches = []
    for row in range(6, payload["last_row"] + 1):
        if cells.get("C" + str(row)):
            current_product, product_cell = _label(cells["C" + str(row)]), "C" + str(row)
        if current_product == product and _label(cells.get("E" + str(row), "")) == _label(destination):
            cell = columns[0] + str(row)
            try:
                value = Decimal(str(cells.get(cell)))
            except InvalidOperation:
                raise ValueError("INVALID_OD_DISTANCE_VALUE") from None
            if not value.is_finite() or value <= 0:
                raise ValueError("INVALID_OD_DISTANCE_VALUE")
            matches.append({"value": str(value), "cell": cell, "product_cell": product_cell,
                            "origin_cell": columns[0] + "5", "destination_cell": "E" + str(row)})
    if len(matches) != 1:
        raise ValueError("OD_DISTANCE_ROUTE_NOT_UNIQUE")
    return matches[0]


@lru_cache(maxsize=2)
def _distance_payload(source: bytes) -> dict:
    workbook = load_workbook(BytesIO(source), read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET]
        cells = {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}
        payload = {"sheet": SHEET, "unit": "NM", "last_row": sheet.max_row, "cells": cells}
    finally:
        workbook.close()
    return payload


def read_distance(path: Path, origin: str, destination: str, product: str) -> dict:
    source = path.read_bytes()
    payload = deepcopy(_distance_payload(source))
    resolved = resolve_distance(payload, origin, destination, product)
    reference = "docs/it_4/" + path.name + " :: " + SHEET + "!" + resolved["cell"]
    return {"version": "OD-DISTANCE.1", **resolved, "unit": "NM", "sheet": SHEET,
            "loading_port": origin, "discharging_port": destination, "cargo_description": product,
            "source_status": "observed_workbook_distance", "source_reference": reference,
            "unit_authority": "User confirms Input - OD Intl values are nautical miles",
            "source_sha256": hashlib.sha256(source).hexdigest(), "raw_payload": payload,
            "payload_sha256": hashlib.sha256(payload_bytes(payload)).hexdigest(),
            "independent_navigation_verified": False}
